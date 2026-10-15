# Genera informes de KDP de prueba (como los que descarga Rodri) a partir de una curva conocida:
# ventas/día = 900 × BSR^(−0,75) en Amazon.es (tapa blanda) y 2.000 × BSR^(−0,8) en Kindle,
# con ruido de Poisson. La prueba de calibración comprueba que el ajuste recupera esa curva.
#   python3 tests/datos/generar-kdp.py   (escribe en tests/fixtures/)
import csv, datetime, math, os, random
from openpyxl import Workbook

random.seed(20261015)
RAIZ = os.path.join(os.path.dirname(__file__), '..', 'fixtures')
INICIO = datetime.date(2026, 8, 1)
DIAS = 61  # agosto y septiembre
LIBROS = [
    # (asin, isbn13, titulo, formato, a, b, bsr_inicial)
    ('B0BR000001', '9798000000011', 'Recuerdos de España', 'impreso', 900, 0.75, 60000),
    ('B0BR000002', '9798000000028', 'Sudoku XXL para mayores', 'impreso', 900, 0.75, 9000),
    ('B0BR000003', '9798000000035', 'Los años 60 y 70', 'impreso', 900, 0.75, 250000),
    ('B0BRKINDL1', None, 'Recuerdos de España (Kindle)', 'kindle', 2000, 0.8, 8000),
]

def poisson(l):
    L, k, p = math.exp(-l), 0, 1.0
    while True:
        p *= random.random()
        if p <= L:
            return k
        k += 1

ventas, lecturas = [], []
for asin, isbn, titulo, formato, a, b, bsr0 in LIBROS:
    bsr = bsr0
    for d in range(DIAS):
        dia = INICIO + datetime.timedelta(days=d)
        # Paseo aleatorio que vuelve hacia su BSR de partida (como un libro con ventas estables).
        bsr = max(300, int(math.exp(math.log(bsr0) + 0.8 * (math.log(bsr) - math.log(bsr0)) + random.gauss(0, 0.35))))
        lam = a * bsr ** (-b)
        u = poisson(lam)
        # Dos lecturas de BSR algunos días (mañana y noche), una el resto; alguno sin lectura.
        if random.random() > 0.12:
            lecturas.append((f'{dia.isoformat()} 10:{random.randint(10, 59)}', asin, bsr, 'amazon.es', 'Tapa blanda' if formato == 'impreso' else 'Versión Kindle'))
            if random.random() < 0.3:
                lecturas.append((f'{dia.isoformat()} 21:{random.randint(10, 59)}', asin, int(bsr * random.uniform(0.85, 1.15)), 'amazon.es', 'Tapa blanda' if formato == 'impreso' else 'Versión Kindle'))
        if u:
            ventas.append((dia, titulo, asin, isbn, formato, u))
        # Ventas en Amazon.com del mismo libro (no deben contar para Amazon.es).
        if random.random() < 0.05:
            ventas.append((dia, titulo, asin, isbn, formato, -1))

# 1) Informe diario en inglés (.xlsx): dos filas de título encima de la cabecera, fechas de Excel.
wb = Workbook()
hojas = {'impreso': wb.active, 'kindle': wb.create_sheet()}
hojas['impreso'].title = 'Paperback Orders'
hojas['kindle'].title = 'eBook Orders'
for ws in hojas.values():
    ws.append(['KDP Orders Report'])
    ws.append(['Sales Period: Aug 1, 2026 - Sep 30, 2026'])
    ws.append(['Date', 'Title', 'Author Name', 'ASIN/ISBN', 'Marketplace', 'Paid Units', 'Free Units'])
for dia, titulo, asin, isbn, formato, u in ventas:
    tienda = 'Amazon.com' if u == -1 else 'Amazon.es'
    ident = (isbn or asin) if formato == 'impreso' else asin
    ws = hojas[formato]
    ws.append([dia, titulo, 'Buen Rato', ident, tienda, 1 if u == -1 else u, 0])
    ws.cell(row=ws.max_row, column=1).number_format = 'yyyy-mm-dd'
wb.save(os.path.join(RAIZ, 'kdp-pedidos-diarios.xlsx'))

# 2) El mismo informe en castellano y en CSV (separador «;», fechas dd/mm/aaaa, ASIN en vez de ISBN).
with open(os.path.join(RAIZ, 'kdp-ventas-castellano.csv'), 'w', newline='', encoding='utf-8-sig') as f:
    w = csv.writer(f, delimiter=';')
    w.writerow(['Fecha', 'Título', 'Nombre del autor', 'ASIN/ISBN', 'Tienda', 'Formato', 'Unidades netas vendidas'])
    for dia, titulo, asin, isbn, formato, u in ventas:
        w.writerow([dia.strftime('%d/%m/%Y'), titulo, 'Buen Rato', asin, 'Amazon.com' if u == -1 else 'Amazon.es', 'Tapa blanda' if formato == 'impreso' else 'eBook', 1 if u == -1 else u])

# 3) Informe mensual (regalías de meses anteriores): una fila por libro y mes.
wb = Workbook()
ws = wb.active
ws.title = 'Combined Sales'
ws.append(['Royalty Date', 'Title', 'Author Name', 'ASIN/ISBN', 'Marketplace', 'Royalty Type', 'Transaction Type', 'Units Sold', 'Units Refunded', 'Net Units Sold'])
mensual = {}
for dia, titulo, asin, isbn, formato, u in ventas:
    if u == -1:
        continue
    k = (dia.strftime('%Y-%m'), titulo, asin, isbn, formato)
    mensual[k] = mensual.get(k, 0) + u
for (mes, titulo, asin, isbn, formato), u in sorted(mensual.items()):
    ws.append([mes, titulo, 'Buen Rato', (isbn or asin) if formato == 'impreso' else asin, 'Amazon.es', '60%' if formato == 'impreso' else '70%', 'Standard', u, 0, u])
ws.append(['2026-09', 'Recuerdos de España (Kindle)', 'Buen Rato', 'B0BRKINDL1', 'Amazon.es', '70%', 'Free - Promotion', 25, 0, 25])
wb.save(os.path.join(RAIZ, 'kdp-regalias-mensual.xlsx'))

# 4) Lecturas de BSR de los libros de Rodri (fecha y hora; ASIN; BSR; tienda; formato).
with open(os.path.join(RAIZ, 'lecturas-bsr.csv'), 'w', newline='', encoding='utf-8-sig') as f:
    w = csv.writer(f, delimiter=';')
    w.writerow(['Fecha', 'ASIN', 'BSR', 'Tienda', 'Formato'])
    for fila in lecturas:
        w.writerow(fila)
print(len(ventas), 'filas de ventas ·', len(lecturas), 'lecturas de BSR')
